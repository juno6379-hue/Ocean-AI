# 파일 역할: 보고서 형식별 본문·표를 추출하고 절·관측소·이슈 단위로 나눕니다.
"""Extract report paragraphs/table rows without inventing PDF pages or dates."""
import re
import struct
import zipfile
import zlib
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Optional
from xml.etree import ElementTree as ET

SUPPORTED = {".pdf", ".txt", ".hwp", ".hwpx", ".xlsx", ".xls", ".docx"}
TYPES = [(('가이드북', '품질관리 가이드'), 'QUALITY_GUIDEBOOK'),
         (('일일상황', '일일현황'), 'DAILY_SITUATION_REPORT'),
         (('일일점검', '일일 점검'), 'DAILY_INSPECTION_REPORT'),
         (('품질처리',), 'QUALITY_PROCESSING_REPORT'),
         (('수집률',), 'QUALITY_COLLECTION_REPORT'),
         (('주간조위', '주간 조위'), 'WEEKLY_TIDE_RESIDUAL_REPORT'),
         (('대조기',), 'SPRING_TIDE_MONITORING_REPORT')]


@dataclass
class Unit:
    text: str
    section: str = "본문"
    page: Optional[int] = None
    locator: str = ""
    context: str = ""


def classify(path):
    for words, kind in TYPES:
        if any(w in str(path) for w in words):
            return kind
    return "OTHER"


def parse_date(text):
    patterns = [r'(?<!\d)((?:19|20)\d{2})\s*[.년/_-]\s*(\d{1,2})\s*[.월/_-]\s*(\d{1,2})(?!\d)',
                r'(?<!\d)((?:19|20)\d{2})(\d{2})(\d{2})(?!\d)']
    for pattern in patterns:
        for m in re.finditer(pattern, text):
            try:
                return datetime(*map(int, m.groups()))
            except ValueError:
                continue
    return None


def report_date(path, units):
    # Filename/nearest parent date outranks old dates mentioned in issue histories.
    for part in [path.stem] + [p.name for p in list(path.parents)[:4]]:
        date = parse_date(part)
        if date:
            return date, "filename_or_parent"
        match = re.search(r'(?<!\d)(2\d)(\d{2})(\d{2})(?!\d)', part)
        if match:
            try:
                return datetime(2000+int(match[1]),int(match[2]),int(match[3])), "filename_yymmdd"
            except ValueError:
                pass
    text = '\n'.join(u.text for u in units[:30])
    for line in text.splitlines():
        if '기준' in line and ('해양관측' in line or '<' in line):
            expanded = re.sub(r"[’‘'](\d{2})\s*\.", r'20\1.', line)
            date = parse_date(expanded)
            if date:
                return date, "report_header"
        if re.search(r'점검일\s*[:：]|보고일\s*[:：]|작성일\s*[:：]|보고일자\s*[:：]', line):
            date = parse_date(line)
            if date:
                return date, "report_header"
    return None, "unresolved"


def xml_paragraphs(data):
    root = ET.fromstring(data)
    for p in root.iter():
        if p.tag.rsplit('}', 1)[-1] == 'p':
            def own_text(node):
                values=[]
                for child in node:
                    tag=child.tag.rsplit('}',1)[-1]
                    if tag=='p':continue  # Nested table paragraphs are visited separately.
                    if tag in {'t','text'}:values.append(child.text or '')
                    else:values.extend(own_text(child))
                return values
            text = ''.join(own_text(p))
            if text.strip():
                yield text.strip()


def spreadsheet_units(sheets):
    units = []
    for title, rows in sheets:
        header = []
        station = None
        for number, row in enumerate(rows, 1):
            values = [str(v).strip() if v is not None else '' for v in row]
            if not any(values):
                continue
            # A real table header must name a station/code or an issue/action column.
            if any(v in {'관측소명', '관측소', '코드', '관측항목', '자료이상 상세내용', '조치내용'} for v in values):
                header = values
                station = None
            elif header:
                for idx, h in enumerate(header):
                    if h in {'관측소명', '관측소'} and idx < len(values):
                        if values[idx]:
                            station = values[idx]
                        elif station:
                            values[idx] = station
            fields = []
            for i, value in enumerate(values):
                if value:
                    label = header[i] if i < len(header) and header[i] else f'열{i+1}'
                    fields.append(f'{label}: {value}')
            units.append(Unit(' | '.join(fields), title, None, f'{title}!row:{number}'))
    return units


def parse(path, should_stop=lambda: False):
    path = Path(path)
    ext = path.suffix.lower()
    if ext == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(''):
            raise ValueError('ENCRYPTED_PDF: password required')
        units=[]
        for i,p in enumerate(reader.pages):
            if should_stop():raise InterruptedError('문서 페이지 추출 중 중단 요청')
            content=p.extract_text(extraction_mode='layout',layout_mode_space_vertically=False) or ''
            if not content.strip():
                content=p.extract_text() or ''  # Layout mode cannot decode some legacy fonts.
            if not content.strip():
                stream=p.get_contents()
                operators=[] if stream is None else [op for _,op in stream.operations]
                # Clip-only pages are blank; any text/image/paint operator requires review/OCR.
                paint={b'Tj',b'TJ',b"'",b'"',b'Do',b'BI',b'sh',b'S',b's',b'f',b'F',b'f*',b'B',b'B*',b'b',b'b*'}
                if not any(op in paint for op in operators):
                    continue  # Truly blank page, not an unreadable scan.
                raise ValueError(f'OCR_REQUIRED: page {i+1} has graphics but no extractable text')
            units.append(Unit(content,'본문',i+1,f'page:{i+1}'))
        return units
    if ext == '.txt':
        data = path.read_bytes()
        for encoding in ('utf-8-sig', 'cp949', 'euc-kr'):
            try:
                return [Unit(data.decode(encoding), locator='text')]
            except UnicodeDecodeError:
                continue
        raise ValueError('ENCODING_UNRESOLVED')
    if ext in {'.hwpx', '.docx'}:
        with zipfile.ZipFile(path) as z:
            names = [n for n in z.namelist() if re.fullmatch(r'Contents/section\d+\.xml', n)] if ext == '.hwpx' else ['word/document.xml']
            units = []
            for name in sorted(names, key=lambda s: int(re.search(r'\d+', s).group()) if re.search(r'\d+', s) else 0):
                units.extend(Unit(t, name, None, name + f':paragraph:{i+1}') for i, t in enumerate(xml_paragraphs(z.read(name))))
            return units
    if ext == '.hwp':
        import olefile
        units = []
        with olefile.OleFileIO(path) as ole:
            header = ole.openstream('FileHeader').read()
            flags = struct.unpack_from('<I', header, 36)[0]
            if flags & (2 | 4):
                raise ValueError('PROTECTED_HWP: encrypted/distribution document')
            sections = sorted(p for p in ole.listdir() if p[0] == 'BodyText' and p[-1].startswith('Section'))
            for section in sections:
                data = ole.openstream(section).read()
                if flags & 1:
                    data = zlib.decompress(data, -15)
                pos = 0
                while pos + 4 <= len(data):
                    record = struct.unpack_from('<I', data, pos)[0]; pos += 4
                    tag, size = record & 0x3ff, record >> 20
                    if size == 0xfff:
                        size = struct.unpack_from('<I', data, pos)[0]; pos += 4
                    payload = data[pos:pos+size]; pos += size
                    if tag == 67:
                        # HWP control records occupy eight UTF-16 code units.
                        words = struct.unpack('<' + 'H' * (len(payload)//2), payload[:len(payload)//2*2])
                        chars=[]; i=0
                        while i < len(words):
                            code=words[i]
                            if code in {1,2,3,4,5,6,7,8,9,11,12,14,15,16,17,18,19,20,21,22,23}:
                                chars.append(' '); i+=8
                            else:
                                if code >= 32 or code in {10,13}: chars.append(chr(code))
                                i+=1
                        text=''.join(chars).strip()
                        if text: units.append(Unit(text, section[-1], None, '/'.join(section)+f':record:{pos}'))
        return units
    if ext == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            return spreadsheet_units((s.title, s.iter_rows(values_only=True)) for s in book.worksheets)
        finally:
            book.close()
    if ext == '.xls':
        import xlrd
        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            return spreadsheet_units((s.name, (s.row_values(i) for i in range(s.nrows))) for s in book.sheets())
        finally:
            book.release_resources()
    raise ValueError('UNSUPPORTED_FORMAT: ' + ext)


STATION_CODE = re.compile(r'\b(?:DT|SO|TW|KG|IE|RT|HF|KHB|SF)_\d{3,6}\b', re.I)
HEADING = re.compile(r'^(?:제\s*\d+\s*[장절]|\d+(?:\.\d+)*[.)]\s+|#{1,6}\s+|표\s*\d+|검사규칙|□\s*)')


def semantic_blocks(units, kind):
    """Section/station/issue/table row boundaries. No fixed token windows/overlap.

    Very long units split only at paragraphs, sentences or table-cell boundaries.
    An indivisible over-context unit fails embedding explicitly rather than truncating.
    """
    output=[]
    section='본문'; block=[]; last_page=None; last_locator=''
    def flush():
        nonlocal block
        if block:
            output.append(Unit('\n'.join(block), section, last_page, last_locator))
            block=[]
    for unit in units:
        # PDF 폰트 제어문자는 검색 원문과 PostgreSQL 양쪽에 동일하게 정규화한다.
        unit=Unit(re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', ' ', unit.text),
                  unit.section,unit.page,unit.locator,unit.context)
        if unit.page != last_page or unit.section not in {'본문',section}:
            flush()
            section=unit.section
        last_page=unit.page; last_locator=unit.locator
        if '!row:' in unit.locator:
            flush(); output.append(unit); continue
        for line in unit.text.splitlines():
            line=line.strip()
            if not line:
                flush(); continue
            if HEADING.match(line):
                flush(); section=line[:200]
            elif (STATION_CODE.search(line) or re.match(r'^(?:관측소|장비|문제|이슈|조치)\s*[:：]',line)
                  or re.match(r"^-?\s*[가-힣A-Za-z]+\s*\([’‘']?\d",line)):
                flush()
            block.append(line)
    flush()
    bounded=[]
    for unit in output:
        if len(unit.text) <= 550:
            bounded.append(unit); continue
        parts=re.split(r'\n+|(?<=[.!?。])\s+|\s+\|\s+',unit.text)
        group=[]
        for part in parts:
            if group and len('\n'.join(group+[part])) > 550:
                bounded.append(Unit('\n'.join(group),unit.section,unit.page,unit.locator,unit.text)); group=[]
            if part.strip():group.append(part.strip())
        if group:bounded.append(Unit('\n'.join(group),unit.section,unit.page,unit.locator,unit.text))
    return bounded
